"""Lectura de TRANSITO\\Novedades.xlsx (agregado 2026-09-19 a pedido del
usuario). Hoja1: Fecha (A), Codigo (B), Descripcion (C), Novedad (D),
Estado (E). Se toma codigo + novedad tal cual, sin filtrar por Estado (no
se pidio). Si el archivo no existe o esta vacio, la app sigue funcionando
igual, simplemente ningun articulo tiene novedad -- no es una fuente
bloqueante como stock o ventas."""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pandas as pd

from src.conversion import normalizar_codigo

logger = logging.getLogger(__name__)


class ExcelNovedades:
    def __init__(self, archivo: Path):
        self.archivo = archivo

    def leer(self) -> pd.DataFrame:
        if not self.archivo.exists():
            logger.info("NOVEDADES: no existe %s, se sigue sin novedades", self.archivo)
            return pd.DataFrame(columns=["codigo", "novedad"])

        wb = openpyxl.load_workbook(self.archivo, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        filas = []
        for fila in ws.iter_rows(min_row=2, max_col=4, values_only=True):
            codigo_raw, novedad_raw = fila[1], fila[3]
            if codigo_raw is None or not novedad_raw or not str(novedad_raw).strip():
                continue
            filas.append({
                "codigo": normalizar_codigo(codigo_raw),
                "novedad": str(novedad_raw).strip(),
            })
        wb.close()

        df = pd.DataFrame(filas, columns=["codigo", "novedad"])
        if not df.empty:
            duplicados = df["codigo"].duplicated().sum()
            if duplicados:
                # Mas de una novedad para el mismo codigo: se combinan en
                # una sola celda en vez de perder alguna.
                df = df.groupby("codigo", as_index=False)["novedad"].agg(" | ".join)
        logger.info("NOVEDADES: archivo=%s filas=%d", self.archivo.name, len(df))
        return df
