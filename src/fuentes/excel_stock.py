"""Lectura de la base de stock (CLAUDE.md seccion 4.1, actualizado
2026-09-17 a pedido del usuario).

Base de stock: un unico archivo manual, D:\\01-Usuario\\Downloads\\
stock_bebidas.xlsx, con una pestaña por deposito -- OB, SV, PR, TC (siglas
confirmadas por el usuario, cada pestaña ES un deposito).

Cada pestaña trae VARIAS tablas independientes, una al lado de la otra (no
alineadas por fila -- cada una tiene su propio orden/longitud):
  - 'DISPONIBLE' (columnas A-D: Articulo, Descripcion Articulo, Cantidad,
    Unids) -- el stock en si. Se usa SOLO este bloque para stock, a pedido
    del usuario (hay un bloque 'FISICO' identico al lado, columnas F-I, que
    NO se usa).
  - Venta Prom Pedido (columna K) + Cod/Descripcion/Venta/ABC (columnas
    L-O): un ranking de ventas aparte, con su PROPIA columna de codigo
    (L, "Cod"). Verificado (2026-09-17): los valores de K difieren por
    pestaña para un mismo codigo (son datos reales por deposito, no una
    tabla duplicada) -- pero el codigo de esa fila es el de la columna L,
    NO el de la columna A (son tablas distintas, no hay que asumir que la
    fila N de A es el mismo articulo que la fila N de L). A pedido del
    usuario, esa columna K se carga como 'venta_promedio_bulto' cruzando
    por (codigo=L, deposito=pestaña) contra las filas de stock.

Estructura de filas: fila 1 = subtitulo de bloque (DISPONIBLE/FISICO), fila
2 = encabezados reales, datos desde fila 3.
"""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pandas as pd

from src.conversion import normalizar_codigo
from src.fuentes.base import FuenteStock

logger = logging.getLogger(__name__)

# Pestañas a leer, en el orden en que se procesan (una = un deposito).
HOJAS_DEPOSITO = ["OB", "SV", "PR", "TC"]


class ExcelStock(FuenteStock):
    def __init__(self, archivo: Path):
        self.archivo = archivo

    def _leer_venta_promedio_bulto(self, ws) -> dict[str, float]:
        """Columna K (Venta Prom Pedido) cruzada por columna L (Cod) -- ver
        docstring del modulo, no tiene relacion posicional con la columna A."""
        mapa: dict[str, float] = {}
        for fila in ws.iter_rows(min_row=3, min_col=11, max_col=12, values_only=True):
            venta_prom, codigo_raw = fila
            if codigo_raw is None or venta_prom is None:
                continue
            try:
                mapa[normalizar_codigo(codigo_raw)] = float(venta_prom)
            except (ValueError, TypeError):
                continue
        return mapa

    def leer(self) -> pd.DataFrame:
        if not self.archivo.exists():
            raise FileNotFoundError(f"No se encontro el archivo de stock: {self.archivo}")
        wb = openpyxl.load_workbook(self.archivo, data_only=True, read_only=True)
        filas = []
        for nombre_hoja in HOJAS_DEPOSITO:
            if nombre_hoja not in wb.sheetnames:
                logger.warning(
                    "STOCK: no se encontro la pestaña '%s' en %s, se omite",
                    nombre_hoja, self.archivo.name,
                )
                continue
            ws = wb[nombre_hoja]
            venta_promedio_bulto = self._leer_venta_promedio_bulto(ws)
            leidas = 0
            con_venta_promedio = 0
            for codigo, descripcion, cantidad, *_resto in ws.iter_rows(min_row=3, max_col=4, values_only=True):
                if codigo is None:
                    continue
                codigo_norm = normalizar_codigo(codigo)
                vpb = venta_promedio_bulto.get(codigo_norm)
                if vpb is not None:
                    con_venta_promedio += 1
                filas.append({
                    "codigo": codigo_norm,
                    "descripcion": (descripcion or "").strip() if isinstance(descripcion, str) else descripcion,
                    "stock_unidades": float(cantidad) if cantidad is not None else 0.0,
                    "deposito": nombre_hoja,
                    "venta_promedio_bulto": vpb,
                })
                leidas += 1
            logger.info(
                "STOCK: pestaña=%s filas_leidas=%d con_venta_promedio_bulto=%d",
                nombre_hoja, leidas, con_venta_promedio,
            )
        wb.close()
        df = pd.DataFrame(filas, columns=["codigo", "descripcion", "stock_unidades", "deposito", "venta_promedio_bulto"])
        logger.info(
            "STOCK: archivo=%s filas_totales=%d depositos=%s",
            self.archivo.name, len(df), sorted(df["deposito"].unique()) if not df.empty else [],
        )
        return df
