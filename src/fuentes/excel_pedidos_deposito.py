"""Lectura de los pedidos pendientes por empresa en la carpeta de Drive
'Tablero' (ej. '1. 1994 - Pedidos.xlsx', '2. Total Center - Pedidos.xlsx',
'4. Kerá Guazú - Pedidos.xlsx') -- reemplaza la lectura anterior de
TRANSITO\\<OB/SV/PR/TC> - Pedidos.xlsx (agregado 2026-09-19), que usaba
codigos de 2 letras que dejaron de existir cuando el stock paso a nombres
completos de deposito (ver excel_stock.py, 2026-10-14).

A diferencia de la fuente anterior, el deposito NO sale del nombre del
archivo (cada archivo es por EMPRESA, no por deposito: una empresa puede
cubrir varios depositos -- ver CRUCE-depositos.jpeg, cruce Deposito<->
Empresa confirmado por el usuario 2026-10-14) sino de la propia fila,
columna 'Depósito' de la pestaña 'Base'. Se filtra contra la misma lista
blanca de 7 depositos reales que usa excel_stock.py (DEPOSITOS_REALES).

Pestaña "Base" (busqueda case-insensitive del nombre): columnas por
posicion D=Remito, G=Depósito, I=Cod., Q=Pedido a retirar (Q ya descuenta
Baja/Agregado/Ajuste CMQ/bloqueos, es la cantidad real pendiente).
CRITERIO (ajustado 2026-10-19 a pedido de los operadores): una fila es
transito solo si la columna D (Remito) esta vacia Y la columna Q es mayor
a 0. Ya NO se usa K=Pedido Original como respaldo cuando Q viene vacia:
Q vacia o <= 0 significa que no hay nada a retirar (ej. pedido bloqueado
por N/D, o ajustado a 0) y no suma transito.
Si el remito ya tiene numero, el pedido ya se desapcho/facturo y no cuenta
como transito pendiente (se omite, tal como pidio el usuario en la fuente
anterior).

NOTA: al 2026-10-14 no existe en la carpeta 'Tablero' un archivo de
pedidos para 'Total Litoral S.A.' (la empresa de Paraná y Victoria segun
CRUCE-depositos.jpeg) -- esos dos depositos quedan sin transito hasta que
aparezca ese archivo. No es un error: la app sigue funcionando con
transito=0 para ellos.

CACHE (agregado 2026-10-15 a pedido del usuario): cada archivo pesa
~1 MB y se reprocesa en cada actualizacion aunque no haya cambiado. Se
guarda en salida\\cache_pedidos_tablero.json, por archivo, su (mtime,
tamaño) y las filas ya parseadas -- si en la proxima corrida ese archivo
tiene el mismo mtime+tamaño, se reusan las filas cacheadas sin volver a
abrir el .xlsx; si cambio (o es nuevo), se reprocesa y se actualiza la
entrada. Un archivo que desaparece de la carpeta se saca del cache."""
from __future__ import annotations

import json
import logging
from pathlib import Path

import openpyxl
import pandas as pd

from src.conversion import normalizar_codigo, sin_acentos
from src.fuentes.excel_stock import DEPOSITOS_REALES

logger = logging.getLogger(__name__)

COL_REMITO = 3  # D (0-based)
COL_DEPOSITO = 6  # G
COL_CODIGO = 8  # I
COL_PEDIDO_A_RETIRAR = 16  # Q

NOMBRE_CACHE = "cache_pedidos_tablero.json"
# Subir este numero cada vez que cambie el criterio de parseo: invalida las
# filas cacheadas aunque el archivo no haya cambiado.
VERSION_CRITERIO = 2  # 2 = solo D vacia y Q > 0 (2026-10-19)


class TransitoPedidosDeposito:
    def __init__(self, carpeta: Path, carpeta_cache: Path | None = None):
        self.carpeta = carpeta
        self.ruta_cache = (carpeta_cache / NOMBRE_CACHE) if carpeta_cache else None

    def _archivos(self) -> list[Path]:
        if not self.carpeta.exists():
            return []
        return sorted(
            a for a in self.carpeta.glob("*.xlsx")
            if not a.name.startswith("~$") and "pedidos" in a.stem.lower()
        )

    def _cargar_cache(self) -> dict:
        if not self.ruta_cache or not self.ruta_cache.exists():
            return {}
        try:
            return json.loads(self.ruta_cache.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            logger.warning("PEDIDOS: no se pudo leer el cache %s (%s), se reprocesa todo", self.ruta_cache, exc)
            return {}

    def _guardar_cache(self, cache: dict) -> None:
        if not self.ruta_cache:
            return
        self.ruta_cache.parent.mkdir(parents=True, exist_ok=True)
        self.ruta_cache.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    def _parsear_archivo(self, archivo: Path) -> list[dict]:
        wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
        nombre_hoja = next((h for h in wb.sheetnames if h.strip().lower() == "base"), None)
        if nombre_hoja is None:
            logger.warning("PEDIDOS: %s no tiene pestaña 'Base', se omite", archivo.name)
            wb.close()
            return []
        ws = wb[nombre_hoja]
        filas_archivo = []
        leidas = 0
        pendientes = 0
        depositos_desconocidos = set()
        for fila in ws.iter_rows(min_row=2, max_col=COL_PEDIDO_A_RETIRAR + 1, values_only=True):
            codigo_raw = fila[COL_CODIGO] if len(fila) > COL_CODIGO else None
            if codigo_raw is None:
                continue
            leidas += 1
            deposito_raw = fila[COL_DEPOSITO] if len(fila) > COL_DEPOSITO else None
            if not deposito_raw:
                continue
            deposito = sin_acentos(str(deposito_raw).strip())
            if deposito not in DEPOSITOS_REALES:
                depositos_desconocidos.add(deposito)
                continue
            remito = fila[COL_REMITO] if len(fila) > COL_REMITO else None
            if remito is not None and str(remito).strip() != "":
                continue  # ya tiene remito -> ya se despacho, no es transito pendiente
            cantidad = fila[COL_PEDIDO_A_RETIRAR] if len(fila) > COL_PEDIDO_A_RETIRAR else None
            try:
                cantidad = float(cantidad)
            except (TypeError, ValueError):
                continue  # Q vacia o no numerica -> nada a retirar
            if cantidad <= 0:
                continue  # Q en 0 (bloqueado/ajustado) -> no es transito
            filas_archivo.append({
                "codigo": normalizar_codigo(codigo_raw),
                "deposito": deposito,
                "transito_bultos_pedidos": cantidad,
            })
            pendientes += 1
        wb.close()
        if depositos_desconocidos:
            logger.info(
                "PEDIDOS: archivo=%s depositos ajenos a la lista de 7 reales, se omiten: %s",
                archivo.name, sorted(depositos_desconocidos),
            )
        logger.info(
            "PEDIDOS: archivo=%s filas=%d pendientes_sin_remito=%d",
            archivo.name, leidas, pendientes,
        )
        return filas_archivo

    def leer(self) -> pd.DataFrame:
        cache = self._cargar_cache()
        cache_nuevo: dict = {}
        filas: list[dict] = []
        reprocesados = 0
        reusados = 0

        for archivo in self._archivos():
            st = archivo.stat()
            huella = {"mtime": st.st_mtime, "size": st.st_size, "version": VERSION_CRITERIO}
            entrada_previa = cache.get(archivo.name)
            if entrada_previa and all(entrada_previa.get(k) == v for k, v in huella.items()):
                filas_archivo = entrada_previa["filas"]
                reusados += 1
                logger.info("PEDIDOS: archivo=%s sin cambios, se reusa el cache (%d filas)", archivo.name, len(filas_archivo))
            else:
                filas_archivo = self._parsear_archivo(archivo)
                reprocesados += 1
            cache_nuevo[archivo.name] = {**huella, "filas": filas_archivo}
            filas.extend(filas_archivo)

        if reprocesados or reusados:
            logger.info(
                "PEDIDOS: %d archivo(s) reprocesados, %d sin cambios (cache reusado)",
                reprocesados, reusados,
            )
        self._guardar_cache(cache_nuevo)

        df = pd.DataFrame(filas, columns=["codigo", "deposito", "transito_bultos_pedidos"])
        if not df.empty:
            df = df.groupby(["codigo", "deposito"], as_index=False)["transito_bultos_pedidos"].sum()
        return df
