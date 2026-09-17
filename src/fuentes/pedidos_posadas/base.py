"""Lectura de archivos de pedido por proveedor (Posadas/TRP).

Copia deliberadamente acotada de la logica ya validada en
D:\\AUTOMATIZACION_PEDIDOS_CONTROL\\automatizacion_pedidos\\extractors\\base.py
(mismo patron leer_pedido_generico, columnas confirmadas contra archivos
reales). Se copia en vez de importarse en vivo de esa otra carpeta para que
UNIFICADOR_STOCK no dependa en tiempo de ejecucion de que ese otro proyecto
siga existiendo o sin cambios -- es de solo lectura de tramites ya resueltos,
no hace falta compartir codigo en caliente. Si esas columnas cambian alguna
vez (verificado contra archivos reales), hay que actualizar ambos lados a
mano.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl


@dataclass
class LineaPedidoTransito:
    codigo: str
    cantidad_unidades: float
    proveedor: str
    fecha_pedido: date
    archivo_origen: str
    # "St Seg" / "D St": piso minimo de dias de cobertura que el proveedor/
    # comprador considera aceptable para ESE producto puntual. None si el
    # archivo no trae esa columna (caso La Serenisima). Se usa unicamente
    # para el color de riesgo (rojo/amarillo/verde/sobrestock) -- nunca se
    # muestra como columna en la web.
    dias_stock_seguridad: float | None = None


def leer_pedido_generico(
    path: Path,
    sheet: str,
    header_row: int,
    col_codigo: int,
    col_cantidad: int,
    proveedor: str,
    col_uxb: int | None = None,
    fecha_cell: tuple[int, int] | None = None,
    fecha_fallback: date | None = None,
    max_col: int = 30,
    col_stock_seguridad: int | None = None,
) -> list[LineaPedidoTransito]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Hoja '{sheet}' no encontrada en {path.name}")
    ws = wb[sheet]

    fecha = None
    if fecha_cell:
        valor = ws.cell(row=fecha_cell[0], column=fecha_cell[1]).value
        if isinstance(valor, datetime):
            fecha = valor.date()
    if fecha is None:
        fecha = fecha_fallback
    if fecha is None:
        raise ValueError(f"No se pudo determinar la fecha del pedido en {path.name}")

    lineas: list[LineaPedidoTransito] = []
    for row in ws.iter_rows(min_row=header_row + 1, max_col=max_col, values_only=True):
        if len(row) < col_codigo or len(row) < col_cantidad:
            continue

        codigo_raw = row[col_codigo - 1]
        cantidad_raw = row[col_cantidad - 1]
        if codigo_raw is None or cantidad_raw in (None, "", " "):
            continue

        try:
            codigo = int(codigo_raw)
            cantidad = float(cantidad_raw)
        except (TypeError, ValueError):
            continue
        if cantidad <= 0:
            continue

        if col_uxb:
            uxb_raw = row[col_uxb - 1] if len(row) >= col_uxb else None
            try:
                uxb = float(uxb_raw)
            except (TypeError, ValueError):
                continue
            if uxb <= 0:
                continue
            cantidad_unidades = cantidad * uxb
        else:
            cantidad_unidades = cantidad

        dias_stock_seguridad = None
        if col_stock_seguridad and len(row) >= col_stock_seguridad:
            valor_seguridad = row[col_stock_seguridad - 1]
            try:
                dias_stock_seguridad = float(valor_seguridad)
            except (TypeError, ValueError):
                dias_stock_seguridad = None
            if dias_stock_seguridad is not None and dias_stock_seguridad <= 0:
                dias_stock_seguridad = None

        lineas.append(
            LineaPedidoTransito(
                codigo=str(codigo),
                cantidad_unidades=cantidad_unidades,
                proveedor=proveedor,
                fecha_pedido=fecha,
                archivo_origen=str(path),
                dias_stock_seguridad=dias_stock_seguridad,
            )
        )

    wb.close()
    return lineas


def parse_fecha_filename(path: Path) -> date | None:
    import re

    match = re.search(r"(\d{1,2})[-.](\d{1,2})[.\-](\d{4})", path.stem)
    if not match:
        return None
    dia, mes, anio = (int(g) for g in match.groups())
    try:
        return date(anio, mes, dia)
    except ValueError:
        return None
